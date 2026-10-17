import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Adjust margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

# Main Title
title = doc.add_heading(level=0)
title_run = title.add_run('सोलर इंस्टॉलेशन फीडबैक फॉर्म')
title_run.font.name = 'Arial'
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(14, 60, 48)  # Deep Forest Green
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
sub_run = subtitle.add_run('क्रिशननुजा रेन्यूएबल्स प्राइवेट लिमिटेड • Arpit Solar Shop')
sub_run.font.size = Pt(11)
sub_run.font.bold = True
sub_run.font.color.rgb = RGBColor(27, 84, 69)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

intro = doc.add_paragraph(
    'क्रिशननुजा रेन्यूएबल्स प्राइवेट लिमिटेड को चुनने के लिए धन्यवाद। हम आपकी प्रतिक्रिया को महत्व देते हैं और सर्वोत्तम सौर ऊर्जा समाधान प्रदान करने का प्रयास करते हैं। कृपया नीचे दिए गए चरणों के अनुसार अपने अनुभव का मूल्यांकन करें (1 से 5 स्टार)।'
)
intro.paragraph_format.space_after = Pt(10)

# Customer details Box
p = doc.add_paragraph()
p.add_run('ग्राहक का नाम (Customer Name): ').bold = True
p.add_run('___________________________    ')
p.add_run('मोबाइल नंबर (Mobile): ').bold = True
p.add_run('___________________\n')
p.add_run('सोलर क्षमता (Capacity): ').bold = True
p.add_run('_______ किलोवाट (kW)    ')
p.add_run('स्थान (Location): ').bold = True
p.add_run('___________________________\n')
p.add_run('दिनांक (Date): ').bold = True
p.add_run('___________________________    ')
p.add_run('सत्यापन (Verification): ').bold = True
p.add_run('[  ] Verified')

doc.add_heading('मूल्यांकन पैमाना (Rating Scale):', level=2)
scale_p = doc.add_paragraph(
    '1 = खराब (Poor)  |  2 = ठीक-ठाक (Fair)  |  3 = अच्छा (Good)  |  4 = बहुत अच्छा (Very Good)  |  5 = उत्कृष्ट (Excellent)'
)
scale_p.paragraph_format.space_after = Pt(10)

# Table for stages
table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'इंस्टॉलेशन के चरण (Installation Stages)'
hdr_cells[1].text = '1 ⭐'
hdr_cells[2].text = '2 ⭐'
hdr_cells[3].text = '3 ⭐'
hdr_cells[4].text = '4 ⭐'
hdr_cells[5].text = '5 ⭐'

for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True

stages = [
    'चरण 1: प्रारंभिक परामर्श और साइट सर्वे (Consultation & Site Survey)',
    'चरण 2: प्रस्ताव, डिजाइन और सिस्टम विवरण (Proposal & System Design)',
    'चरण 3: टीम का पेशेवर व्यवहार और समय पर इंस्टॉलेशन (Installation Team & Timeliness)',
    'चरण 4: नेट मीटरिंग और सरकारी सब्सिडी (PM Surya Ghar Subsidy) प्रक्रिया में सहायता',
    'चरण 5: इंस्टॉलेशन के बाद साइट की सफाई और हैंडओवर (Site Cleanup & Handover)',
    'चरण 6: सिस्टम का प्रदर्शन और बिजली उत्पादन (System Performance & Generation)',
    'चरण 7: क्रिशननुजा रेन्यूएबल्स के साथ कुल मिलाकर संतुष्टि (Overall Satisfaction)',
    'चरण 8: आप अपने दोस्तों और परिवार को Arpit Solar Shop की सोलर सेवाएँ recommend करने के लिए कितनी स्टार रेटिंग देंगे? (Recommendation)'
]

for item in stages:
    row_cells = table.add_row().cells
    row_cells[0].text = item
    for i in range(1, 6):
        row_cells[i].text = '[   ]'

doc.add_paragraph()
doc.add_heading('अतिरिक्त सुझाव या टिप्पणियां (Additional Comments / Suggestions):', level=2)
doc.add_paragraph('_________________________________________________________________________________\n_________________________________________________________________________________')

doc.add_paragraph('\nग्राहक के हस्ताक्षर (Customer Signature): ___________________________    दिनांक: ___________________')

# Google Review Section
doc.add_paragraph()
google_box = doc.add_paragraph()
g_run1 = google_box.add_run('⭐ Google पर हमें 5-स्टार रेटिंग देकर अपना अनुभव साझा करें!\n')
g_run1.bold = True
g_run1.font.color.rgb = RGBColor(180, 83, 9)
g_run2 = google_box.add_run('कृपया 30 सेकंड का समय निकालकर नीचे दिए गए Google लिंक पर 5-स्टार रिव्यू दें:\n')
g_run3 = google_box.add_run('Google Review Link: https://g.page/r/CcJZzDtcNJTgEAE/review\n')
g_run3.bold = True
g_run3.font.color.rgb = RGBColor(27, 84, 69)

doc.save('solar_feedback_form_hindi.docx')
print('Hindi docx created successfully with Google Review link: solar_feedback_form_hindi.docx')
